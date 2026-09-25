"""DOCX and PDF rendering of output versions (FR-OUT-006, ADR-024).

Both formats carry the same content as the Markdown and HTML exports:
- the title and blocks;
- labels and traces according to mode;
- protected quotes copied verbatim, numbered in referenced modes;
- the reference list built from records;
- the integrity footer.

Arabic is shaped and laid out right to left. Qur'an quotes use a Qur'an font.
Quote text goes to both renderers unchanged:
- A DOCX stores text as Unicode, so every quote is read back from the finished file before it is
  returned. A quote that did not survive exactly is never shipped.
- A PDF stores shaped glyphs. It also embeds `quotes.json`, the exact quote texts with their SHA-256,
  so a reader or a test can check them by machine.
"""

from __future__ import annotations

import hashlib
import io
import json
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from uuid import UUID

import docx
from docx.document import Document as DocxFile
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph
from fpdf import FPDF
from sqlalchemy.orm import Session

from research_api.config import get_settings
from research_api.contracts.enums import OutputBlockKind, OutputMode, QuoteSourceKind
from research_api.modules.outputs_integrity import render
from research_api.modules.outputs_integrity.render import INTEGRITY, REFERENCES, Document

K = OutputBlockKind
FONTS = Path(__file__).resolve().parents[2] / "assets" / "fonts"
_RTL = re.compile(r"[֐-ࣿיִ-﷿ﹰ-﻿]")


class RenderingError(Exception):
    """A protected quote did not survive rendering; the file is not returned."""


@dataclass(frozen=True)
class Element:
    kind: str  # title | heading | quote | list | claim | text | references | footer
    text: str = ""
    level: int = 2
    items: list[str] = field(default_factory=list)
    note: str = ""  # quote marker, label and trace, shown small and muted
    quran: bool = False


def _rtl(text: str, default: bool = False) -> bool:
    """Whether the first strong character is right-to-left; text without one follows the document."""
    for char in text:
        direction = unicodedata.bidirectional(char)
        if direction in ("R", "AL"):
            return True
        if direction == "L":
            return False
    return default


def elements(session: Session, project_id: UUID, doc: Document) -> list[Element]:
    """The document laid out once, in the same order and with the same rules as the Markdown export."""
    referenced = doc.mode is not OutputMode.READABLE
    audit = doc.mode is OutputMode.AUDIT
    out = [Element("title", doc.title)]
    number = 0
    for block in doc.blocks:
        label = f"[{block.label}]" if referenced and block.label else ""
        trace = ", ".join(f"{t.entity_type}:{t.entity_id}" for t in block.trace) if audit and block.trace else ""
        if block.kind is K.HEADING:
            out.append(Element("heading", block.text, level=min(block.level or 2, 4)))
        elif block.kind is K.QUOTE and block.quote is not None:
            number += 1
            note = " ".join(p for p in (f"[{number}]" if referenced else "", label, trace) if p)
            quran = block.quote.source_kind is QuoteSourceKind.QURAN
            out.append(Element("quote", block.quote.text, note=note, quran=quran))
        elif block.kind is K.LIST:
            out.append(Element("list", items=list(block.items or []), note=" ".join(p for p in (label, trace) if p)))
        else:
            kind = "claim" if block.kind is K.CLAIM else "text"
            out.append(Element(kind, block.text, note=" ".join(p for p in (label, trace) if p)))
    citations = render._citations(session, project_id, doc) if referenced else []
    if citations:
        out.append(Element("references", REFERENCES.get(doc.language, REFERENCES["en"]), items=citations))
    if render._footer(doc):
        footer = f"{INTEGRITY.get(doc.language, 'Integrity')}: {doc.integrity} · v{doc.version} {doc.status}"
        out.append(Element("footer", footer))
    return out


def _quotes(items: list[Element]) -> list[str]:
    return [e.text for e in items if e.kind == "quote"]


# --- DOCX ---


def _bidi(paragraph: Paragraph) -> None:
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:bidi"))


def _run(
    paragraph: Paragraph,
    text: str,
    *,
    bold: bool = False,
    size: float | None = None,
    muted: bool = False,
    quran: bool = False,
) -> None:
    run = paragraph.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if muted:
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    font = "Amiri Quran" if quran else "Amiri"
    fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:cs"), font)
    if quran:
        fonts.set(qn("w:ascii"), font)
        fonts.set(qn("w:hAnsi"), font)
    if _RTL.search(text):
        run._r.get_or_add_rPr().append(OxmlElement("w:rtl"))


class _Docx:
    def __init__(self, doc: Document) -> None:
        self.rtl_doc = doc.language == "ar"
        self.file: DocxFile = docx.Document()
        normal = self.file.styles["Normal"]
        normal.font.name = "Amiri"
        normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), "Amiri")
        self.file.core_properties.title = doc.title
        self.file.core_properties.language = doc.language

    def paragraph(self, text: str, style: str | None = None, *, heading: int | None = None) -> Paragraph:
        paragraph = (
            self.file.add_heading(level=heading) if heading is not None else self.file.add_paragraph(style=style)
        )
        if _rtl(text, self.rtl_doc):
            _bidi(paragraph)
        return paragraph

    def note(self, element: Element) -> None:
        if element.note:
            _run(self.paragraph(element.note), element.note, size=8, muted=True)

    def add(self, element: Element) -> None:
        if element.kind == "title":
            _run(self.paragraph(element.text, heading=0), element.text)
        elif element.kind in ("heading", "references"):
            _run(self.paragraph(element.text, heading=element.level), element.text)
            for number, citation in enumerate(element.items, start=1):
                _run(self.paragraph(citation), f"{number}. {citation}")
        elif element.kind == "quote":
            for line in element.text.split("\n"):
                _run(self.paragraph(line, "Quote"), line, size=16 if element.quran else None, quran=element.quran)
            self.note(element)
        elif element.kind == "list":
            for item in element.items:
                _run(self.paragraph(item, "List Bullet"), item)
            self.note(element)
        else:
            paragraph = self.paragraph(element.text)
            muted = element.kind == "footer"
            _run(paragraph, element.text, bold=element.kind == "claim", muted=muted, size=9 if muted else None)
            if element.note:
                _run(paragraph, f" {element.note}", size=8, muted=True)


def docx_bytes(session: Session, project_id: UUID, doc: Document) -> bytes:
    items = elements(session, project_id, doc)
    writer = _Docx(doc)
    for element in items:
        writer.add(element)
    buffer = io.BytesIO()
    writer.file.save(buffer)
    data = buffer.getvalue()
    _check_docx(data, _quotes(items))
    return data


def _check_docx(data: bytes, quotes: list[str]) -> None:
    paragraphs = {p.text for p in docx.Document(io.BytesIO(data)).paragraphs}
    for quote in quotes:
        if any(line not in paragraphs for line in quote.split("\n")):
            raise RenderingError("a protected quote did not survive DOCX rendering")


# --- PDF ---


def _quran_font() -> Path:
    configured = get_settings().export_quran_font_path
    return configured if configured is not None and configured.is_file() else FONTS / "AmiriQuran.ttf"


MARKER = 7  # mm reserved for a list number or bullet


class _Pdf(FPDF):
    def __init__(self, rtl_doc: bool) -> None:
        super().__init__(format="A4")
        self.rtl_doc = rtl_doc
        self.set_auto_page_break(auto=True, margin=18)
        self.add_font("Amiri", "", str(FONTS / "Amiri-Regular.ttf"))
        self.add_font("Amiri", "B", str(FONTS / "Amiri-Bold.ttf"))
        self.add_font("Quran", "", str(_quran_font()))
        self.set_text_shaping(True)
        self.add_page()

    def block(
        self,
        text: str,
        *,
        size: float,
        style: str = "",
        family: str = "Amiri",
        indent: float = 0,
        muted: bool = False,
        gap: float = 2,
        marker: str = "",
        rtl: bool | None = None,
    ) -> None:
        """A paragraph aligned by its own direction. A list marker is drawn on the leading side as its own cell,
        because the line layout would otherwise put a leading number or bullet on the left of Arabic text."""
        self.set_font(family, style, size)
        self.set_text_color(0x66 if muted else 0)
        right = _rtl(text, self.rtl_doc) if rtl is None else rtl
        self.set_text_shaping(True, direction="rtl" if right else "ltr")
        height = size * 0.5
        width = self.epw - indent - (MARKER if marker else 0)
        if marker:
            top = self.get_y()
            self.set_x(self.w - self.r_margin - indent - MARKER if right else self.l_margin + indent)
            self.cell(MARKER, height, marker, align="R" if right else "L")
            self.set_y(top)
        self.set_x(self.l_margin if right else self.l_margin + indent + (MARKER if marker else 0))
        self.multi_cell(width, height, text, align="R" if right else "L", new_x="LMARGIN", new_y="NEXT")
        self.ln(gap)

    def note(self, element: Element, indent: float = 0) -> None:
        """Marker, label and trace follow the direction of the block they annotate."""
        if element.note:
            owner = element.text or (element.items[0] if element.items else "")
            self.block(element.note, size=8, indent=indent, muted=True, gap=0, rtl=_rtl(owner, self.rtl_doc))

    def quote(self, element: Element) -> None:
        top = self.get_y()
        family, size = ("Quran", 16) if element.quran else ("Amiri", 12)
        for line in element.text.split("\n"):
            self.block(line, size=size, family=family, indent=8, gap=0)
        x = self.w - self.r_margin - 3 if _rtl(element.text, self.rtl_doc) else self.l_margin + 3
        self.set_draw_color(0x88)
        self.line(x, top, x, self.get_y())
        self.note(element, indent=8)
        self.ln(3)

    def add(self, element: Element) -> None:
        if element.kind == "title":
            self.block(element.text, size=20, style="B", gap=4)
        elif element.kind in ("heading", "references"):
            self.block(element.text, size={1: 18, 2: 15, 3: 13}.get(element.level, 12), style="B")
            for number, citation in enumerate(element.items, start=1):
                self.block(citation, size=10, gap=1, marker=f"{number}.")
        elif element.kind == "quote":
            self.quote(element)
        elif element.kind == "list":
            for item in element.items:
                self.block(item, size=12, indent=4, gap=0, marker="•")
            self.note(element)
            self.ln(2)
        elif element.kind == "footer":
            self.set_draw_color(0x88)
            self.line(self.l_margin, self.get_y() + 2, self.w - self.r_margin, self.get_y() + 2)
            self.ln(4)
            self.block(element.text, size=9, muted=True)
        else:
            self.block(element.text, size=12, style="B" if element.kind == "claim" else "", gap=0)
            self.note(element)
            self.ln(2)


def pdf_bytes(session: Session, project_id: UUID, doc: Document) -> bytes:
    items = elements(session, project_id, doc)
    pdf = _Pdf(doc.language == "ar")
    pdf.set_title(doc.title)
    pdf.set_lang(doc.language)
    pdf.set_creator("Integrated AI Research System")
    for element in items:
        pdf.add(element)
    manifest = [
        {"number": i, "text": q, "sha256": hashlib.sha256(q.encode()).hexdigest()}
        for i, q in enumerate(_quotes(items), start=1)
    ]
    pdf.embed_file(
        basename="quotes.json",
        bytes=json.dumps(manifest, ensure_ascii=False, indent=2).encode(),
        desc="Exact text of every protected quote in this document",
    )
    return bytes(pdf.output())
