"""DOCX and PDF export of output versions (issue #52, FR-OUT-006, ADR-024)."""

from __future__ import annotations

import hashlib
import io
import json
from typing import Any
from uuid import UUID

import docx
import pytest
from fastapi.testclient import TestClient
from pypdf import PdfReader
from pytest import MonkeyPatch
from sqlalchemy.orm import Session

from research_api.contracts.enums import OutputMode
from research_api.modules.outputs_integrity import office, render
from research_api.modules.outputs_integrity.schemas import Block
from research_api.platform import queue
from tests.integration.test_evidence_hypotheses import _accept, _excerpt
from tests.integration.test_outputs import PASSAGE, _create, _project_with_evidence
from tests.integration.test_reference import approved_quran  # noqa: F401 - fixture

ARABIC = "يتراجع التوجيه بشدة في السنة الثانية."
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture(autouse=True)
def _no_broker(monkeypatch: MonkeyPatch) -> None:
    monkeypatch.setattr(queue, "dispatch", lambda task, job_id: None)


def _url(pid: str, output: dict[str, Any]) -> str:
    return f"/api/v1/projects/{pid}/outputs/{output['id']}/versions/{output['latest']['id']}/export"


def _paragraphs(data: bytes) -> list[Any]:
    return list(docx.Document(io.BytesIO(data)).paragraphs)


def _quotes_json(data: bytes) -> list[dict[str, Any]]:
    reader = PdfReader(io.BytesIO(data))
    [attachment] = reader.attachments["quotes.json"]
    return list(json.loads(attachment))


def _arabic_project(client: TestClient) -> str:
    pid = str(
        client.post("/api/v1/projects", json={"title": "ع", "initial_input": "x", "input_type": "IDEA"}).json()["id"]
    )
    claim = client.post(
        f"/api/v1/projects/{pid}/claims", json={"statement": "التوجيه يتراجع", "claim_type": "OBSERVATION"}
    ).json()
    _, excerpt = _excerpt(client, pid, "دراسة الفوج", ARABIC)
    _accept(client, pid, ("CLAIM", claim["id"]), "SUPPORTS", excerpt, "SUPPORTED")
    return pid


def test_docx_carries_quotes_references_and_integrity(client: TestClient) -> None:
    pid, _, _ = _project_with_evidence(client)
    output = _create(client, pid, mode="AUDIT")
    response = client.get(_url(pid, output), params={"format": "docx"})
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == DOCX
    assert response.headers["content-disposition"].endswith('.docx"')
    texts = [p.text for p in _paragraphs(response.content)]
    assert texts[0] == "Year-two report"
    assert PASSAGE in texts, "the quote is its own paragraph, verbatim"
    assert any("Cohort study" in t for t in texts), "references are built from records"
    assert any(t.startswith("Integrity: NOT_CHECKED") for t in texts)
    assert any("Claim:" in t for t in texts), "audit mode shows traces"


def test_arabic_docx_is_right_to_left(client: TestClient) -> None:
    pid = _arabic_project(client)
    output = _create(client, pid, language="ar")
    data = client.get(_url(pid, output), params={"format": "docx"}).content
    [quote] = [p for p in _paragraphs(data) if p.text == ARABIC]
    assert quote._p.pPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bidi") is not None
    assert any("المراجع" in p.text for p in _paragraphs(data))


def test_pdf_embeds_exact_quotes_and_renders_them(client: TestClient) -> None:
    pid, _, _ = _project_with_evidence(client)
    output = _create(client, pid)
    response = client.get(_url(pid, output), params={"format": "pdf"})
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")
    assert _quotes_json(response.content) == [
        {"number": 1, "text": PASSAGE, "sha256": hashlib.sha256(PASSAGE.encode()).hexdigest()}
    ]
    text = PdfReader(io.BytesIO(response.content)).pages[0].extract_text()
    assert PASSAGE in text and "Year-two report" in text


def test_arabic_pdf_shapes_the_quote_without_changing_it(client: TestClient) -> None:
    pid = _arabic_project(client)
    output = _create(client, pid, language="ar")
    data = client.get(_url(pid, output), params={"format": "pdf"}).content
    assert _quotes_json(data)[0]["text"] == ARABIC
    reader = PdfReader(io.BytesIO(data))
    # Extraction reorders neutral punctuation at the end of right-to-left lines; quotes.json is the exact record.
    assert ARABIC.rstrip(".") in "".join(page.extract_text() for page in reader.pages)
    assert reader.metadata is not None and reader.metadata.title == "Year-two report"


def test_quran_quotes_use_the_quran_font(client: TestClient, approved_quran: dict[str, Any]) -> None:  # noqa: F811
    pid = str(
        client.post("/api/v1/projects", json={"title": "Q", "initial_input": "x", "input_type": "IDEA"}).json()["id"]
    )
    hid = client.post(f"/api/v1/projects/{pid}/hypotheses", json={"content": {"statement": "h"}}).json()["id"]
    review = client.post(
        f"/api/v1/projects/{pid}/reference-reviews",
        json={
            "target_type": "HYPOTHESIS",
            "target_id": hid,
            "question": "Q?",
            "analytical_category": "VALUES_AND_EVALUATIVE_STANDARDS",
        },
    ).json()
    client.post(
        f"/api/v1/projects/{pid}/reference-reviews/{review['id']}/entries",
        json={"layer": "SOURCE_TEXT", "quran_ref": "1:2"},
    )
    output = _create(client, pid, output_type="REFERENCE_REVIEW", language="ar")
    pdf = client.get(_url(pid, output), params={"format": "pdf"}).content
    assert [q["text"] for q in _quotes_json(pdf)] == ["PLACEHOLDER-V1-1-2"]
    fonts = {
        str(font["/BaseFont"])
        for page in PdfReader(io.BytesIO(pdf)).pages
        for font in page["/Resources"]["/Font"].values()  # type: ignore[index]
    }
    assert any("AmiriQuran" in f for f in fonts), fonts
    paragraphs = _paragraphs(client.get(_url(pid, output), params={"format": "docx"}).content)
    [quote] = [p for p in paragraphs if p.text == "PLACEHOLDER-V1-1-2"]
    assert (
        quote.runs[0]._r.rPr.rFonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cs")
        == "Amiri Quran"
    )


def test_a_quote_that_does_not_survive_is_never_shipped(session: Session, client: TestClient) -> None:
    pid, _, _ = _project_with_evidence(client)
    output = _create(client, pid)
    blocks = [Block.model_validate(b) for b in output["latest"]["blocks"]]
    doc = render.Document("Report", "en", OutputMode.REFERENCED, blocks, 1, "DRAFT")
    data = office.docx_bytes(session, UUID(pid), doc)
    with pytest.raises(office.RenderingError):
        office._check_docx(data, [PASSAGE + " (altered)"])
